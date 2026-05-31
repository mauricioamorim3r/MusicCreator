from __future__ import annotations

import base64
import csv
import io
import json
import mimetypes
from datetime import datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

from services.config import CACHE_DIR, ensure_runtime_dirs
from services.copilot_audio import build_audio_attachment_dossier


ATTACHMENTS_DIR = CACHE_DIR / "copilot_attachments"
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
TEXT_EXTENSIONS = {".txt", ".md", ".json", ".csv", ".log", ".yaml", ".yml"}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".xlsx"}
AUDIO_EXTENSIONS = {".mp3", ".wav", ".flac", ".m4a", ".aac", ".ogg"}
SUPPORTED_EXTENSIONS = IMAGE_EXTENSIONS | TEXT_EXTENSIONS | DOCUMENT_EXTENSIONS | AUDIO_EXTENSIONS
MAX_ATTACHMENT_BYTES = 15 * 1024 * 1024
MAX_AUDIO_ATTACHMENT_BYTES = 200 * 1024 * 1024
MAX_EXTRACTED_CHARS = 40_000


def capture_screen() -> dict[str, Any]:
    ensure_runtime_dirs()
    ATTACHMENTS_DIR.mkdir(parents=True, exist_ok=True)
    try:
        from PIL import ImageGrab

        image = ImageGrab.grab(all_screens=True)
    except Exception as exc:
        raise RuntimeError(
            "Não foi possível capturar a tela neste ambiente. "
            "Use o upload de imagem como alternativa."
        ) from exc

    path = ATTACHMENTS_DIR / f"screen_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid4().hex[:6]}.png"
    image.save(path, format="PNG")
    return prepare_attachment(path.name, path.read_bytes())


def prepare_attachment(
    name: str,
    content: bytes,
    *,
    enable_audio_premium: bool = False,
) -> dict[str, Any]:
    if not content:
        raise ValueError("O arquivo anexado está vazio.")
    suffix = Path(name).suffix.lower()
    size_limit = MAX_AUDIO_ATTACHMENT_BYTES if suffix in AUDIO_EXTENSIONS else MAX_ATTACHMENT_BYTES
    if len(content) > size_limit:
        if suffix in AUDIO_EXTENSIONS:
            raise ValueError("O áudio anexado excede o limite de 200 MB.")
        raise ValueError("O arquivo anexado excede o limite de 15 MB.")

    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            "Formato não suportado. Use áudio, imagem, TXT, Markdown, JSON, CSV, PDF, DOCX ou XLSX."
        )

    mime_type = mimetypes.guess_type(name)[0] or "application/octet-stream"
    kind = "image" if suffix in IMAGE_EXTENSIONS else "audio" if suffix in AUDIO_EXTENSIONS else "document"
    attachment: dict[str, Any] = {
        "name": Path(name).name,
        "suffix": suffix,
        "mime_type": mime_type,
        "size_bytes": len(content),
        "kind": kind,
    }
    if suffix in IMAGE_EXTENSIONS:
        attachment["image_base64"] = base64.b64encode(content).decode("ascii")
        attachment["extracted_text"] = ""
    elif suffix in AUDIO_EXTENSIONS:
        dossier = build_audio_attachment_dossier(
            name,
            content,
            enable_premium=enable_audio_premium,
        )
        attachment["audio_path"] = str(dossier.pop("_audio_path", ""))
        attachment["audio_analysis"] = dossier
        attachment["extracted_text"] = ""
    else:
        attachment["extracted_text"] = _extract_document_text(suffix, content)
    return attachment


def public_attachment_summary(attachment: dict[str, Any]) -> dict[str, Any]:
    return {
        "name": attachment.get("name"),
        "kind": attachment.get("kind"),
        "mime_type": attachment.get("mime_type"),
        "size_bytes": attachment.get("size_bytes"),
        "extracted_text": attachment.get("extracted_text", ""),
        "image_sent_to_llm": bool(attachment.get("image_base64")),
        "audio_analysis": attachment.get("audio_analysis"),
    }


def _extract_document_text(suffix: str, content: bytes) -> str:
    if suffix in {".txt", ".md", ".log", ".yaml", ".yml"}:
        return _truncate(content.decode("utf-8", errors="replace"))
    if suffix == ".json":
        try:
            parsed = json.loads(content.decode("utf-8", errors="replace"))
            return _truncate(json.dumps(parsed, ensure_ascii=False, indent=2))
        except Exception:
            return _truncate(content.decode("utf-8", errors="replace"))
    if suffix == ".csv":
        text = content.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(text))
        rows = [", ".join(row) for row in list(reader)[:250]]
        return _truncate("\n".join(rows))
    if suffix == ".pdf":
        return _extract_pdf(content)
    if suffix == ".docx":
        return _extract_docx(content)
    if suffix == ".xlsx":
        return _extract_xlsx(content)
    return ""


def _extract_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        pages = []
        for index, page in enumerate(reader.pages[:40], start=1):
            pages.append(f"## Página {index}\n{page.extract_text() or ''}")
        return _truncate("\n\n".join(pages))
    except Exception as exc:
        return f"[PDF recebido, mas o texto não pôde ser extraído localmente: {exc}]"


def _extract_docx(content: bytes) -> str:
    try:
        from docx import Document

        document = Document(io.BytesIO(content))
        return _truncate("\n".join(paragraph.text for paragraph in document.paragraphs))
    except Exception as exc:
        return f"[DOCX recebido, mas o texto não pôde ser extraído localmente: {exc}]"


def _extract_xlsx(content: bytes) -> str:
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        lines: list[str] = []
        for worksheet in workbook.worksheets[:5]:
            lines.append(f"## Planilha: {worksheet.title}")
            for row in worksheet.iter_rows(max_row=150, values_only=True):
                lines.append(" | ".join("" if value is None else str(value) for value in row[:20]))
        return _truncate("\n".join(lines))
    except Exception as exc:
        return f"[XLSX recebido, mas o conteúdo não pôde ser extraído localmente: {exc}]"


def _truncate(text: str) -> str:
    if len(text) <= MAX_EXTRACTED_CHARS:
        return text
    return text[:MAX_EXTRACTED_CHARS] + "\n[conteúdo truncado localmente]"
